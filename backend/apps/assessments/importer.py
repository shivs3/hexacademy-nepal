"""Bulk question import (CSV or XLSX) and export (CSV).

Columns (header row required, names case-insensitive):

    subject, topic, kind, difficulty, question, option_a ... option_h,
    correct, explanation, marks, negative_marks, tags, status

- subject: name or slug of an existing subject; topic: name or slug of an
  existing topic in it (optional).
- kind: single | multiple | true_false (default single). For true_false the
  options default to "True" and "False" and `correct` is true or false.
- correct: option letters, e.g. "B" or "A,C".
- tags: separated by ";".
- Text is plain text: it is HTML-escaped and wrapped in paragraphs. Formulas
  written as \\( ... \\) or $$ ... $$ become maths nodes.

The import is all or nothing: if any row has an error nothing is saved, and
the report lists every error with its row number.
"""

import csv
import html
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook

from apps.audit import services as audit
from common.exceptions import DomainError

from . import services
from .models import Difficulty, Question, Subject, Topic

MAX_ROWS = 2000
MAX_BYTES = 2 * 1024 * 1024
LETTERS = "abcdefgh"
OPTION_COLUMNS = [f"option_{c}" for c in LETTERS]
COLUMNS = [
    "subject",
    "topic",
    "kind",
    "difficulty",
    "question",
    *OPTION_COLUMNS,
    "correct",
    "explanation",
    "marks",
    "negative_marks",
    "tags",
    "status",
]
TRUE_WORDS = {"true", "t", "yes", "1"}
FALSE_WORDS = {"false", "f", "no", "0"}

_BLOCK_MATH = re.compile(r"\$\$(.+?)\$\$", re.S)
_INLINE_MATH = re.compile(r"\\\((.+?)\\\)", re.S)


def text_to_html(text: str) -> str:
    """Plain text -> safe paragraphs with maths nodes. The result is sanitised again by the services."""
    text = (text or "").strip()
    if not text:
        return ""

    def block(m):
        return f'</p><div data-type="block-math" data-latex="{html.escape(m.group(1).strip(), quote=True)}"></div><p>'

    def inline(m):
        return f'<span data-type="inline-math" data-latex="{html.escape(m.group(1).strip(), quote=True)}"></span>'

    paragraphs = []
    for para in re.split(r"\n\s*\n", text):
        # Escape everything except the maths markers, then turn markers into nodes.
        pieces, last = [], 0
        for m in re.finditer(r"\$\$.+?\$\$|\\\(.+?\\\)", para, re.S):
            pieces.append(html.escape(para[last : m.start()]).replace("\n", "<br>"))
            token = m.group(0)
            pieces.append(_BLOCK_MATH.sub(block, token) if token.startswith("$$") else _INLINE_MATH.sub(inline, token))
            last = m.end()
        pieces.append(html.escape(para[last:]).replace("\n", "<br>"))
        paragraphs.append(f"<p>{''.join(pieces)}</p>")
    return "".join(paragraphs).replace("<p></p>", "")


def html_to_text(value: str) -> str:
    """Inverse used by export: maths nodes back to markers, tags stripped."""
    value = re.sub(
        r'<div data-type="block-math" data-latex="([^"]*)"></div>', lambda m: f"$${html.unescape(m.group(1))}$$", value
    )
    value = re.sub(
        r'<span data-type="inline-math" data-latex="([^"]*)"></span>',
        lambda m: f"\\({html.unescape(m.group(1))}\\)",
        value,
    )
    value = re.sub(r"<br\s*/?>", "\n", value)
    value = re.sub(r"</p>\s*<p>", "\n\n", value)
    return html.unescape(re.sub(r"<[^>]+>", "", value)).strip()


# --- Reading -------------------------------------------------------------------------


def read_rows(upload) -> list[dict]:
    if upload.size > MAX_BYTES:
        raise DomainError("The file is larger than 2 MB.", code="file_too_large")
    name = (upload.name or "").lower()
    raw = upload.read()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for malformed files
            raise DomainError("The spreadsheet could not be read.", code="invalid_file") from exc
        sheet = workbook.worksheets[0]
        values = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DomainError("Save the CSV file as UTF-8.", code="invalid_file") from exc
        values = list(csv.reader(io.StringIO(text)))
    else:
        raise DomainError("Upload a .csv or .xlsx file.", code="invalid_file")
    if not values:
        raise DomainError("The file is empty.", code="invalid_file")
    header = [h.strip().lower().replace(" ", "_") for h in values[0]]
    missing = [c for c in ("subject", "question", "correct") if c not in header]
    if missing:
        raise DomainError(f"Missing column(s): {', '.join(missing)}.", code="invalid_file")
    rows = [dict(zip(header, row, strict=False)) for row in values[1:] if any(cell.strip() for cell in row)]
    if len(rows) > MAX_ROWS:
        raise DomainError(f"Import at most {MAX_ROWS} questions at a time.", code="too_many_rows")
    return rows


# --- Parsing -------------------------------------------------------------------------


class RowError(Exception):
    def __init__(self, field, message):
        super().__init__(message)
        self.field, self.message = field, message


def _decimal(value: str, field: str, default: Decimal) -> Decimal:
    if not (value or "").strip():
        return default
    try:
        number = Decimal(value.strip())
    except InvalidOperation as exc:
        raise RowError(field, "Must be a number.") from exc
    if number < 0 or number > 1000 or number != number.quantize(Decimal("0.01")):
        raise RowError(field, "Must be between 0 and 1000 with at most two decimals.")
    return number


class Lookup:
    def __init__(self):
        self.subjects = {}
        for s in Subject.objects.all():
            self.subjects[s.slug.lower()] = s
            self.subjects[s.name.lower()] = s
        self.topics = {}
        for t in Topic.objects.all():
            self.topics[(t.subject_id, t.slug.lower())] = t
            self.topics[(t.subject_id, t.name.lower())] = t


# "\\(" means the file was written with every backslash doubled (JSON or code
# string escaping). Imported as is, each formula keeps stray backslashes and
# fails to render, so reject the row instead.
_DOUBLED_MARKER = re.compile(r"\\\\[()]")


def _check_maths(field: str, text: str):
    if _DOUBLED_MARKER.search(text):
        raise RowError(field, r"The maths has doubled backslashes. Write \( \frac{a}{b} \), not \\( \\frac{a}{b} \\).")


def parse_row(row: dict, lookup: Lookup) -> tuple[dict, list[dict]]:
    get = lambda key: (row.get(key) or "").strip()  # noqa: E731
    _check_maths("question", get("question"))
    _check_maths("explanation", get("explanation"))
    for column in OPTION_COLUMNS:
        _check_maths("options", get(column))
    subject = lookup.subjects.get(get("subject").lower())
    if subject is None:
        raise RowError("subject", f"No subject called “{get('subject')}”.")
    topic = None
    if get("topic"):
        topic = lookup.topics.get((subject.pk, get("topic").lower()))
        if topic is None:
            raise RowError("topic", f"No topic called “{get('topic')}” in {subject.name}.")
    kind = (get("kind") or "single").lower().replace("-", "_").replace(" ", "_")
    if kind not in Question.Kind.values:
        raise RowError("kind", "Use single, multiple or true_false.")
    difficulty = (get("difficulty") or "medium").lower()
    if difficulty not in Difficulty.values:
        raise RowError("difficulty", "Use easy, medium or hard.")
    status = (get("status") or "draft").lower()
    if status not in Question.Status.values:
        raise RowError("status", "Use draft, active or retired.")
    prompt = text_to_html(get("question"))
    if not prompt:
        raise RowError("question", "The question is empty.")

    texts = [get(c) for c in OPTION_COLUMNS]
    correct_raw = get("correct").lower()
    if kind == Question.Kind.TRUE_FALSE:
        if not any(texts):
            texts = ["True", "False"]
        if correct_raw in TRUE_WORDS:
            correct_raw = "a"
        elif correct_raw in FALSE_WORDS:
            correct_raw = "b"
    while texts and not texts[-1]:
        texts.pop()
    if any(not t for t in texts):
        raise RowError("options", "Fill the options in order without gaps.")
    letters = {c.strip() for c in re.split(r"[,;\s]+", correct_raw) if c.strip()}
    if not letters or any(len(c) != 1 or c not in LETTERS[: len(texts)] for c in letters):
        raise RowError("correct", f"Use option letters A–{LETTERS[max(len(texts), 1) - 1].upper()}.")
    options = [{"text": text_to_html(t), "is_correct": LETTERS[i] in letters} for i, t in enumerate(texts)]
    try:
        services.validate_options(kind, options)
    except DomainError as exc:
        raise RowError("options", str(exc.detail)) from exc

    tags = sorted({t.strip().lower()[:40] for t in get("tags").split(";") if t.strip()})[:20]
    data = {
        "subject": subject,
        "topic": topic,
        "kind": kind,
        "difficulty": difficulty,
        "status": status,
        "prompt": prompt,
        "explanation": text_to_html(get("explanation")),
        "default_marks": _decimal(get("marks"), "marks", Decimal("1")),
        "default_negative_marks": _decimal(get("negative_marks"), "negative_marks", Decimal("0")),
        "tags": tags,
    }
    if data["default_marks"] <= 0:
        raise RowError("marks", "Marks must be more than zero.")
    return data, options


def import_questions(user, upload, *, course=None, dry_run=False, request=None) -> dict:
    rows = read_rows(upload)
    lookup = Lookup()
    parsed, errors = [], []
    for number, row in enumerate(rows, start=2):  # row 1 is the header
        try:
            parsed.append(parse_row(row, lookup))
        except RowError as exc:
            errors.append({"row": number, "field": exc.field, "message": exc.message})
    report = {"rows": len(rows), "valid": len(parsed), "created": 0, "errors": errors, "dry_run": dry_run}
    if errors or dry_run or not parsed:
        return report
    with transaction.atomic():
        for data, options in parsed:
            services.create_question(user, {**data, "course": course}, options, request=request, audit_event=False)
    report["created"] = len(parsed)
    audit.record(
        "assessments.questions_imported",
        actor=user,
        request=request,
        count=len(parsed),
        course=course.pk if course else None,
    )
    return report


# --- Export --------------------------------------------------------------------------


def export_csv(questions) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for q in questions.select_related("subject", "topic").prefetch_related("options"):
        options = list(q.options.all())
        texts = [html_to_text(o.text) for o in options] + [""] * (len(LETTERS) - len(options))
        correct = ",".join(LETTERS[i].upper() for i, o in enumerate(options) if o.is_correct)
        writer.writerow(
            [
                q.subject.slug,
                q.topic.slug if q.topic else "",
                q.kind,
                q.difficulty,
                html_to_text(q.prompt),
                *texts[: len(LETTERS)],
                correct,
                html_to_text(q.explanation),
                q.default_marks,
                q.default_negative_marks,
                ";".join(q.tags),
                q.status,
            ]
        )
    return buffer.getvalue()


def bank_filter(qs, params):
    """Shared by the list and export endpoints."""
    if subject := params.get("subject"):
        qs = qs.filter(subject_id=subject) if str(subject).isdigit() else qs.none()
    if topic := params.get("topic"):
        if str(topic).isdigit():
            qs = qs.filter(Q(topic_id=topic) | Q(topic__parent_id=topic))
    for field in ("kind", "status", "difficulty"):
        if value := params.get(field):
            qs = qs.filter(**{field: value})
    if course := params.get("course"):
        if course == "shared":
            qs = qs.filter(course__isnull=True)
        elif str(course).isdigit():
            qs = qs.filter(course_id=course)
    if tag := params.get("tag"):
        qs = qs.filter(tags__contains=[tag])
    if search := (params.get("search") or "").strip():
        match = Q(prompt__icontains=search) | Q(tags__contains=[search.lower()])
        if search.isdigit():
            match |= Q(pk=int(search))
        qs = qs.filter(match)
    return qs
